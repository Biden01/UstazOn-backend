"""
Presentation Generation Service for creating PowerPoint presentations
"""
import logging
import json
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

logger = logging.getLogger(__name__)


class PresentationService:
    """Service for generating PowerPoint presentations"""

    @staticmethod
    def create_presentation(presentation_data: dict) -> BytesIO:
        """
        Create a PowerPoint presentation from structured data

        Args:
            presentation_data: Dictionary with presentation structure:
                {
                    "title": "Presentation Title",
                    "subject": "Subject Name",
                    "grade": "5 класс",
                    "slides": [
                        {
                            "type": "title",
                            "title": "Main Title",
                            "subtitle": "Subtitle"
                        },
                        {
                            "type": "content",
                            "title": "Slide Title",
                            "content": ["Point 1", "Point 2", "Point 3"]
                        },
                        {
                            "type": "two_column",
                            "title": "Title",
                            "left_content": ["Left 1", "Left 2"],
                            "right_content": ["Right 1", "Right 2"]
                        }
                    ]
                }

        Returns:
            BytesIO object containing the PPTX file
        """
        try:
            prs = Presentation()
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)

            for slide_data in presentation_data.get("slides", []):
                slide_type = slide_data.get("type", "content")

                if slide_type == "title":
                    PresentationService._add_title_slide(prs, slide_data)
                elif slide_type == "content":
                    PresentationService._add_content_slide(prs, slide_data)
                elif slide_type == "two_column":
                    PresentationService._add_two_column_slide(prs, slide_data)
                elif slide_type == "image":
                    PresentationService._add_image_slide(prs, slide_data)
                else:
                    PresentationService._add_content_slide(prs, slide_data)

            # Save to BytesIO
            pptx_output = BytesIO()
            prs.save(pptx_output)
            pptx_output.seek(0)

            return pptx_output

        except Exception as e:
            logger.error(f"Error creating presentation: {e}")
            raise

    @staticmethod
    def _add_title_slide(prs: Presentation, slide_data: dict):
        """Add title slide"""
        slide_layout = prs.slide_layouts[0]  # Title slide layout
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = slide_data.get("title", "")
        subtitle.text = slide_data.get("subtitle", "")

        # Style title
        title.text_frame.paragraphs[0].font.size = Pt(44)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)

        # Style subtitle
        subtitle.text_frame.paragraphs[0].font.size = Pt(28)
        subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(89, 89, 89)

    @staticmethod
    def _add_content_slide(prs: Presentation, slide_data: dict):
        """Add content slide with bullet points"""
        slide_layout = prs.slide_layouts[1]  # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        content = slide.placeholders[1]

        title.text = slide_data.get("title", "")

        # Add content
        text_frame = content.text_frame
        text_frame.clear()

        content_items = slide_data.get("content", [])
        if isinstance(content_items, str):
            content_items = [content_items]

        for idx, item in enumerate(content_items):
            if idx == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()

            p.text = item
            p.level = 0
            p.font.size = Pt(18)

        # Style title
        title.text_frame.paragraphs[0].font.size = Pt(32)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)

    @staticmethod
    def _add_two_column_slide(prs: Presentation, slide_data: dict):
        """Add slide with two columns"""
        slide_layout = prs.slide_layouts[3]  # Two content layout
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        title.text = slide_data.get("title", "")

        # Left column
        left_content = slide.placeholders[1]
        left_text_frame = left_content.text_frame
        left_text_frame.clear()

        left_items = slide_data.get("left_content", [])
        for idx, item in enumerate(left_items):
            if idx == 0:
                p = left_text_frame.paragraphs[0]
            else:
                p = left_text_frame.add_paragraph()
            p.text = item
            p.font.size = Pt(16)

        # Right column
        right_content = slide.placeholders[2]
        right_text_frame = right_content.text_frame
        right_text_frame.clear()

        right_items = slide_data.get("right_content", [])
        for idx, item in enumerate(right_items):
            if idx == 0:
                p = right_text_frame.paragraphs[0]
            else:
                p = right_text_frame.add_paragraph()
            p.text = item
            p.font.size = Pt(16)

        # Style title
        title.text_frame.paragraphs[0].font.size = Pt(32)
        title.text_frame.paragraphs[0].font.bold = True

    @staticmethod
    def _add_image_slide(prs: Presentation, slide_data: dict):
        """Add slide with high-quality image from Unsplash/Pexels or AI-generated"""
        try:
            import os
            import urllib.parse
            import urllib.request
            
            slide_layout = prs.slide_layouts[6]  # Blank layout
            slide = prs.slides.add_slide(slide_layout)

            # Add title
            left = Inches(0.5)
            top = Inches(0.5)
            width = Inches(9)
            height = Inches(1)

            title_box = slide.shapes.add_textbox(left, top, width, height)
            title_frame = title_box.text_frame
            title_frame.text = slide_data.get("title", "")
            if title_frame.paragraphs:
                title_frame.paragraphs[0].font.size = Pt(32)
                title_frame.paragraphs[0].font.bold = True
                title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

            # Get image prompt/query
            image_prompt = slide_data.get("image_prompt", "")
            if not image_prompt:
                image_prompt = slide_data.get('title', 'education')
            
            # Clean up prompt for search query
            search_query = image_prompt.replace('"', '').replace("'", "")
            encoded_query = urllib.parse.quote(search_query)
            
            image_data = None
            
            # Try Unsplash first (best quality, free)
            unsplash_key = os.getenv("UNSPLASH_ACCESS_KEY", "")
            if unsplash_key and not image_data:
                try:
                    unsplash_url = f"https://api.unsplash.com/search/photos?query={encoded_query}&orientation=landscape&per_page=1"
                    req = urllib.request.Request(
                        unsplash_url,
                        headers={
                            'Authorization': f'Client-ID {unsplash_key}',
                            'User-Agent': 'Mozilla/5.0'
                        }
                    )
                    with urllib.request.urlopen(req, timeout=10) as response:
                        data = json.loads(response.read().decode())
                        if data.get("results") and len(data["results"]) > 0:
                            img_url = data["results"][0]["urls"]["regular"]
                            img_req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
                            with urllib.request.urlopen(img_req, timeout=15) as img_response:
                                image_data = BytesIO(img_response.read())
                            logger.info(f"Got image from Unsplash for: {search_query}")
                except Exception as e:
                    logger.warning(f"Unsplash failed: {e}")
            
            # Try Pexels as second option
            pexels_key = os.getenv("PEXELS_API_KEY", "")
            if pexels_key and not image_data:
                try:
                    pexels_url = f"https://api.pexels.com/v1/search?query={encoded_query}&orientation=landscape&per_page=1"
                    req = urllib.request.Request(
                        pexels_url,
                        headers={
                            'Authorization': pexels_key,
                            'User-Agent': 'Mozilla/5.0'
                        }
                    )
                    with urllib.request.urlopen(req, timeout=10) as response:
                        data = json.loads(response.read().decode())
                        if data.get("photos") and len(data["photos"]) > 0:
                            img_url = data["photos"][0]["src"]["large"]
                            img_req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
                            with urllib.request.urlopen(img_req, timeout=15) as img_response:
                                image_data = BytesIO(img_response.read())
                            logger.info(f"Got image from Pexels for: {search_query}")
                except Exception as e:
                    logger.warning(f"Pexels failed: {e}")
            
            # Fallback to Pollinations AI generation
            if not image_data:
                try:
                    # Make prompt more educational/professional
                    ai_prompt = f"professional educational illustration, {image_prompt}, high quality, clean design, suitable for classroom presentation"
                    encoded_prompt = urllib.parse.quote(ai_prompt)
                    image_url = f"https://image.pollinations.ai/prompt/{encoded_prompt}?width=1024&height=768&nologo=true"
                    
                    req = urllib.request.Request(
                        image_url,
                        headers={'User-Agent': 'Mozilla/5.0'}
                    )
                    with urllib.request.urlopen(req, timeout=30) as response:
                        image_data = BytesIO(response.read())
                    logger.info(f"Got image from Pollinations AI for: {search_query}")
                except Exception as e:
                    logger.error(f"Pollinations AI failed: {e}")
            
            # Add image to slide if we got one
            if image_data:
                img_left = Inches(1)
                img_top = Inches(1.8)
                img_width = Inches(8)
                img_height = Inches(5)
                
                slide.shapes.add_picture(image_data, img_left, img_top, width=img_width, height=img_height)
            else:
                # Fallback to placeholder text if all sources fail
                note_box = slide.shapes.add_textbox(Inches(3), Inches(3), Inches(4), Inches(1))
                note_frame = note_box.text_frame
                note_frame.text = "[Изображение недоступно]"
                note_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                
        except Exception as e:
            logger.error(f"Error adding image slide: {e}")


# Create singleton instance
presentation_service = PresentationService()
